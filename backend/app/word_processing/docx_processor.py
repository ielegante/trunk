import io
import logging
import os
import tempfile
from pathlib import Path
from typing import Any, BinaryIO, Dict, List, Optional, Union

import mammoth
from app.word_processing.word_xml_parser import (
    WordDocumentStructure,
    WordElement,
    WordElementType,
    WordXMLParser,
)
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

logger = logging.getLogger(__name__)


class DocxProcessor:
    """Advanced DOCX document processor using python-docx and custom XML parsing"""

    def __init__(self):
        self.xml_parser = WordXMLParser()

    def load_document(self, file_path: str) -> Document:
        """Load a DOCX document from file path"""
        try:
            return Document(file_path)
        except Exception as e:
            logger.error(f"Failed to load document {file_path}: {str(e)}")
            raise

    def load_document_from_bytes(self, docx_bytes: bytes) -> Document:
        """Load a DOCX document from bytes"""
        try:
            return Document(io.BytesIO(docx_bytes))
        except Exception as e:
            logger.error(f"Failed to load document from bytes: {str(e)}")
            raise

    def extract_comprehensive_structure(self, file_path: str) -> WordDocumentStructure:
        """Extract comprehensive document structure using XML parser"""
        return self.xml_parser.parse_docx_file(file_path)

    def extract_comprehensive_structure_from_bytes(
        self, docx_bytes: bytes
    ) -> WordDocumentStructure:
        """Extract comprehensive document structure from bytes"""
        return self.xml_parser.parse_docx_bytes(docx_bytes)

    def extract_text_content(self, document: Document) -> str:
        """Extract plain text content from document"""
        text_parts = []

        for paragraph in document.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Extract text from tables
        for table in document.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    cell_text = " ".join(
                        p.text for p in cell.paragraphs if p.text.strip()
                    )
                    row_text.append(cell_text)
                if any(row_text):
                    text_parts.append(" | ".join(row_text))

        return "\n\n".join(text_parts)

    def convert_to_markdown(self, file_path: str) -> str:
        """Convert DOCX to Markdown using mammoth"""
        try:
            with open(file_path, "rb") as docx_file:
                result = mammoth.convert_to_markdown(docx_file)
                return result.value
        except Exception as e:
            logger.error(f"Failed to convert to markdown: {str(e)}")
            # Fallback to basic conversion
            return self._convert_to_markdown_fallback(file_path)

    def convert_bytes_to_markdown(self, docx_bytes: bytes) -> str:
        """Convert DOCX bytes to Markdown"""
        try:
            result = mammoth.convert_to_markdown(io.BytesIO(docx_bytes))
            return result.value
        except Exception as e:
            logger.error(f"Failed to convert bytes to markdown: {str(e)}")
            # Fallback to basic conversion
            return self._convert_bytes_to_markdown_fallback(docx_bytes)

    def _convert_to_markdown_fallback(self, file_path: str) -> str:
        """Fallback markdown conversion using python-docx"""
        document = self.load_document(file_path)
        return self._document_to_markdown(document)

    def _convert_bytes_to_markdown_fallback(self, docx_bytes: bytes) -> str:
        """Fallback markdown conversion for bytes"""
        document = self.load_document_from_bytes(docx_bytes)
        return self._document_to_markdown(document)

    def _document_to_markdown(self, document: Document) -> str:
        """Convert Document object to Markdown"""
        markdown_parts = []

        for paragraph in document.paragraphs:
            if not paragraph.text.strip():
                continue

            # Determine paragraph type
            style_name = paragraph.style.name.lower()

            if "heading" in style_name:
                # Extract heading level
                level = 1
                if "heading 1" in style_name:
                    level = 1
                elif "heading 2" in style_name:
                    level = 2
                elif "heading 3" in style_name:
                    level = 3
                elif "heading 4" in style_name:
                    level = 4
                elif "heading 5" in style_name:
                    level = 5
                elif "heading 6" in style_name:
                    level = 6

                markdown_parts.append(f"{'#' * level} {paragraph.text}")

            elif "list" in style_name or "bullet" in style_name:
                markdown_parts.append(f"- {paragraph.text}")

            else:
                # Regular paragraph with formatting
                formatted_text = self._format_runs_to_markdown(paragraph.runs)
                if formatted_text.strip():
                    markdown_parts.append(formatted_text)

        # Process tables
        for table in document.tables:
            table_markdown = self._table_to_markdown(table)
            if table_markdown:
                markdown_parts.append(table_markdown)

        return "\n\n".join(markdown_parts)

    def _format_runs_to_markdown(self, runs) -> str:
        """Format text runs to Markdown with formatting"""
        formatted_parts = []

        for run in runs:
            text = run.text
            if not text:
                continue

            # Apply formatting
            if run.bold and run.italic:
                text = f"***{text}***"
            elif run.bold:
                text = f"**{text}**"
            elif run.italic:
                text = f"*{text}*"

            if run.underline:
                text = f"<u>{text}</u>"

            formatted_parts.append(text)

        return "".join(formatted_parts)

    def _table_to_markdown(self, table) -> str:
        """Convert table to Markdown format"""
        if not table.rows:
            return ""

        markdown_rows = []

        for i, row in enumerate(table.rows):
            row_cells = []
            for cell in row.cells:
                cell_text = " ".join(
                    p.text.strip() for p in cell.paragraphs if p.text.strip()
                )
                row_cells.append(cell_text or " ")

            markdown_row = f"| {' | '.join(row_cells)} |"
            markdown_rows.append(markdown_row)

            # Add header separator after first row
            if i == 0:
                separator = f"| {' | '.join(['---'] * len(row_cells))} |"
                markdown_rows.append(separator)

        return "\n".join(markdown_rows)

    def extract_document_metadata(self, document: Document) -> Dict[str, Any]:
        """Extract document metadata"""
        core_props = document.core_properties

        metadata = {
            "title": core_props.title,
            "author": core_props.author,
            "subject": core_props.subject,
            "keywords": core_props.keywords,
            "comments": core_props.comments,
            "category": core_props.category,
            "created": core_props.created.isoformat() if core_props.created else None,
            "modified": (
                core_props.modified.isoformat() if core_props.modified else None
            ),
            "last_modified_by": core_props.last_modified_by,
            "revision": core_props.revision,
            "version": core_props.version,
        }

        # Add document statistics
        metadata.update(
            {
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
                "word_count": self._count_words(document),
                "character_count": self._count_characters(document),
            }
        )

        return metadata

    def _count_words(self, document: Document) -> int:
        """Count words in document"""
        text = self.extract_text_content(document)
        return len(text.split())

    def _count_characters(self, document: Document) -> int:
        """Count characters in document"""
        text = self.extract_text_content(document)
        return len(text)

    def extract_styles(self, document: Document) -> Dict[str, Dict]:
        """Extract style definitions from document"""
        styles = {}

        for style in document.styles:
            if style.type == WD_STYLE_TYPE.PARAGRAPH:
                styles[style.name] = {
                    "type": "paragraph",
                    "built_in": style.builtin,
                    "hidden": style.hidden,
                    "priority": style.priority,
                    "properties": self._extract_style_properties(style),
                }
            elif style.type == WD_STYLE_TYPE.CHARACTER:
                styles[style.name] = {
                    "type": "character",
                    "built_in": style.builtin,
                    "hidden": style.hidden,
                    "priority": style.priority,
                    "properties": self._extract_style_properties(style),
                }

        return styles

    def _extract_style_properties(self, style) -> Dict:
        """Extract properties from a style"""
        properties = {}

        try:
            # Paragraph properties
            if hasattr(style, "paragraph_format"):
                para_format = style.paragraph_format
                properties["paragraph"] = {
                    "alignment": para_format.alignment,
                    "left_indent": para_format.left_indent,
                    "right_indent": para_format.right_indent,
                    "first_line_indent": para_format.first_line_indent,
                    "space_before": para_format.space_before,
                    "space_after": para_format.space_after,
                    "line_spacing": para_format.line_spacing,
                }

            # Font properties
            if hasattr(style, "font"):
                font = style.font
                properties["font"] = {
                    "name": font.name,
                    "size": font.size.pt if font.size else None,
                    "bold": font.bold,
                    "italic": font.italic,
                    "underline": font.underline,
                    "color": (
                        str(font.color.rgb) if font.color and font.color.rgb else None
                    ),
                }
        except Exception as e:
            logger.warning(f"Failed to extract style properties: {str(e)}")

        return properties

    def create_document_from_markdown(self, markdown_content: str) -> Document:
        """Create a new DOCX document from Markdown content"""
        document = Document()

        lines = markdown_content.split("\n")
        current_paragraph = []
        in_code_block = False

        for line in lines:
            line = line.rstrip()

            # Handle code blocks
            if line.startswith("```"):
                if current_paragraph:
                    self._add_paragraph_to_document(
                        document, "\n".join(current_paragraph)
                    )
                    current_paragraph = []
                in_code_block = not in_code_block
                continue

            if in_code_block:
                # Add code line with monospace formatting
                p = document.add_paragraph()
                run = p.add_run(line)
                run.font.name = "Courier New"
                continue

            # Handle headings
            if line.startswith("#"):
                if current_paragraph:
                    self._add_paragraph_to_document(
                        document, "\n".join(current_paragraph)
                    )
                    current_paragraph = []

                level = len(line) - len(line.lstrip("#"))
                heading_text = line.lstrip("#").strip()

                if level <= 6:
                    document.add_heading(heading_text, level)
                else:
                    document.add_paragraph(heading_text)
                continue

            # Handle lists
            if line.startswith("- ") or line.startswith("* "):
                if current_paragraph:
                    self._add_paragraph_to_document(
                        document, "\n".join(current_paragraph)
                    )
                    current_paragraph = []

                list_text = line[2:].strip()
                p = document.add_paragraph(list_text, style="List Bullet")
                continue

            # Handle numbered lists
            if line.strip() and line[0].isdigit() and ". " in line:
                if current_paragraph:
                    self._add_paragraph_to_document(
                        document, "\n".join(current_paragraph)
                    )
                    current_paragraph = []

                list_text = line.split(". ", 1)[1] if ". " in line else line
                p = document.add_paragraph(list_text, style="List Number")
                continue

            # Handle empty lines
            if not line.strip():
                if current_paragraph:
                    self._add_paragraph_to_document(
                        document, "\n".join(current_paragraph)
                    )
                    current_paragraph = []
                continue

            # Regular paragraph content
            current_paragraph.append(line)

        # Add remaining paragraph
        if current_paragraph:
            self._add_paragraph_to_document(document, "\n".join(current_paragraph))

        return document

    def _add_paragraph_to_document(self, document: Document, paragraph_text: str):
        """Add a paragraph with formatting to document"""
        if not paragraph_text.strip():
            return

        paragraph = document.add_paragraph()

        # Parse inline formatting
        parts = self._parse_inline_formatting(paragraph_text)

        for part_text, formatting in parts:
            run = paragraph.add_run(part_text)

            if formatting.get("bold"):
                run.bold = True
            if formatting.get("italic"):
                run.italic = True
            if formatting.get("underline"):
                run.underline = True

    def _parse_inline_formatting(self, text: str) -> List[tuple]:
        """Parse inline formatting in text"""
        # Simplified implementation - would need more sophisticated parsing
        parts = []

        # For now, just return plain text
        parts.append((text, {}))

        return parts

    def save_document(self, document: Document, file_path: str):
        """Save document to file"""
        try:
            document.save(file_path)
        except Exception as e:
            logger.error(f"Failed to save document to {file_path}: {str(e)}")
            raise

    def save_document_to_bytes(self, document: Document) -> bytes:
        """Save document to bytes"""
        try:
            doc_io = io.BytesIO()
            document.save(doc_io)
            doc_io.seek(0)
            return doc_io.getvalue()
        except Exception as e:
            logger.error(f"Failed to save document to bytes: {str(e)}")
            raise

    def compare_documents(self, doc1_path: str, doc2_path: str) -> Dict[str, Any]:
        """Compare two DOCX documents"""
        doc1 = self.load_document(doc1_path)
        doc2 = self.load_document(doc2_path)

        # Extract text content
        text1 = self.extract_text_content(doc1)
        text2 = self.extract_text_content(doc2)

        # Basic comparison
        comparison = {
            "identical": text1 == text2,
            "doc1_stats": {
                "paragraphs": len(doc1.paragraphs),
                "tables": len(doc1.tables),
                "words": len(text1.split()),
                "characters": len(text1),
            },
            "doc2_stats": {
                "paragraphs": len(doc2.paragraphs),
                "tables": len(doc2.tables),
                "words": len(text2.split()),
                "characters": len(text2),
            },
        }

        return comparison
